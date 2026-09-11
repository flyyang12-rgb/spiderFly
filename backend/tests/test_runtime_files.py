from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from spiderfly_runtime import TaskError
from spiderfly_runtime.average import average
from spiderfly_runtime.excel import read_excel
from spiderfly_runtime.excel_write import write_excel
from spiderfly_runtime.files import list_files
from spiderfly_runtime.table_filter import filter_equals


class RuntimeFileTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)

    def book(self, rows):
        path = self.root / "input.xlsx"
        book = Workbook()
        for row in rows:
            book.active.append(row)
        book.create_sheet("说明")["A1"] = "保留"
        try:
            book.save(path)
        finally:
            book.close()
        return path

    def test_round_trip_preserves_zero_leading_zeroes_order_and_source(self):
        path = self.book([["订单号", "状态", "金额"], ["00123", "待处理", 0],
                          ["00124", "已完成", 19.5], ["00125", "待处理", 42.5]])
        original = path.read_bytes()
        table = read_excel(str(path), required_columns=["状态"])
        selected = filter_equals(table.columns, table.rows, "状态", "待处理")
        output = self.root / "result.xlsx"
        write_excel(str(output), selected.columns, selected.rows)
        result = read_excel(str(output))
        self.assertEqual([r["订单号"] for r in result.rows], ["00123", "00125"])
        self.assertEqual([r["金额"] for r in result.rows], [0, 42.5])
        self.assertEqual(path.read_bytes(), original)

    def test_no_matches_keeps_headers_and_no_overwrite(self):
        path = self.book([["状态"], ["已完成"]])
        table = read_excel(str(path))
        result = filter_equals(table.columns, table.rows, "状态", "待处理")
        output = self.root / "empty.xlsx"
        write_excel(str(output), result.columns, result.rows)
        self.assertEqual(read_excel(str(output)).row_count, 0)
        before = output.read_bytes()
        with self.assertRaises(TaskError):
            write_excel(str(output), result.columns, result.rows)
        self.assertEqual(output.read_bytes(), before)

    def test_append_preserves_other_sheets_and_rejects_changed_source_cells(self):
        path = self.book([["订单号", "金额"], ["001", "5,2"]])
        table = read_excel(str(path))
        rows = [{**r, "平均数": average(r["金额"]).average} for r in table.rows]
        output = self.root / "append.xlsx"
        write_excel(str(output), table.columns + ["平均数"], rows,
                    sheet_name=table.sheet_name, template_file=str(path))
        book = load_workbook(output)
        try:
            self.assertEqual(book["说明"]["A1"].value, "保留")
            self.assertEqual(book.active["C2"].value, 3.5)
        finally:
            book.close()
        rows[0]["订单号"] = "changed"
        bad = self.root / "bad.xlsx"
        with self.assertRaises(TaskError):
            write_excel(str(bad), table.columns + ["平均数"], rows,
                        sheet_name=table.sheet_name, template_file=str(path))
        self.assertFalse(bad.exists())

    def test_invalid_headers_and_formulas_are_rejected(self):
        for rows in ([["状态", "状态"], [1, 2]], [["状态"], ["=1+1"]], [[""], [1]]):
            with self.subTest(rows=rows), self.assertRaises(TaskError):
                read_excel(str(self.book(rows)))
        path = self.book([["状态"], ["待处理"]])
        with self.assertRaises(TaskError) as error:
            read_excel(str(path), required_columns=["金额"])
        self.assertEqual(error.exception.code, "EXCEL_COLUMNS_MISSING")

    def test_strict_types_and_nonfinite_numbers(self):
        for value in (True, None, [], "", "1,,2", "nan", float("inf")):
            with self.subTest(value=value), self.assertRaises(TaskError):
                average(value)
        for value, expected in (("5,2", 3.5), (0, 0.0), ("-1,1", 0.0)):
            self.assertEqual(average(value).average, expected)
        rows = [{"key": True}, {"key": 1}, {"key": "1"}, {"key": None}, {"key": ""}]
        self.assertEqual(filter_equals(["key"], rows, "key", 1).rows, [{"key": 1}])
        self.assertEqual(rows[0], {"key": True})

    def test_file_listing_is_nonrecursive_case_insensitive_and_read_only(self):
        (self.root / "a.XLSX").write_bytes(b"unchanged")
        (self.root / "b.txt").write_text("unchanged")
        (self.root / "nested").mkdir()
        (self.root / "nested/hidden.xlsx").write_bytes(b"unchanged")
        self.assertEqual(list_files(str(self.root), "*.xlsx").files, [str(self.root / "a.XLSX")])
        self.assertEqual((self.root / "a.XLSX").read_bytes(), b"unchanged")
        with self.assertRaises(TaskError):
            list_files(str(self.root), "../*")

    def test_incorrect_output_is_not_reported_as_verified(self):
        with patch("spiderfly_runtime.average._average", return_value={"average": 999.0, "count": 2}):
            with self.assertRaises(TaskError) as error:
                average("1,3")
        self.assertEqual(error.exception.code, "VERIFICATION_FAILED")

    def test_installed_task_entry_needs_no_server_or_table_imports(self):
        code = """
import importlib.abc, sys
class Boundary(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'app', 'pydantic', 'openpyxl', 'spiderfly_instructions', 'example_flows'}:
            raise ImportError(fullname)
sys.meta_path.insert(0, Boundary())
from spiderfly_runtime import TaskContext, TaskResult, run_task
assert TaskResult('ok').code == 'TASK_DONE'
"""
        result = subprocess.run([sys.executable, "-I", "-c", code], cwd=self.root,
                                capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_release_contains_only_runtime(self):
        wheel = Path(__file__).resolve().parents[2] / "release/runtime/spiderfly_runtime-0.1.0-py3-none-any.whl"
        with ZipFile(wheel) as archive:
            names = archive.namelist()
            self.assertTrue(any(n == "spiderfly_runtime/task.py" for n in names))
            self.assertFalse(any(n.startswith(("example_flows/", "spiderfly_instructions/", "app/", "flows/")) for n in names))


if __name__ == "__main__":
    unittest.main()
