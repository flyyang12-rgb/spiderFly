"""Read-only execution profile and external result checks, never controlled by the model."""
from __future__ import annotations
import asyncio
import ast
import base64
from collections import Counter
from difflib import SequenceMatcher
import csv
import io
import json
import os
from pathlib import Path
import re
import subprocess
import zipfile
from urllib.parse import urlsplit

from . import ai_tools

DRIVER = Path(__file__).resolve().parents[1] / 'sandbox/driver.py'
SUPPORTED = {'requests': '2.32.5', 'openpyxl': '3.1.5', 'beautifulsoup4': '4.13.5'}
MAX_OUTPUT = 12 * 1024 * 1024

def safe_name(name):
    reserved = {'CON','PRN','AUX','NUL',*(f'COM{x}' for x in range(10)),*(f'LPT{x}' for x in range(10))}
    return isinstance(name,str) and bool(re.fullmatch(r'[\w.-]{1,100}',name)) and not name.endswith('.') and name.split('.')[0].upper() not in reserved

def command(check=False):
    path = DRIVER.resolve().as_posix()
    if os.name == 'nt':
        path = '/mnt/' + path[0].lower() + path[2:]
        args = ['wsl.exe', '-d', 'Ubuntu', '--', 'python3', path]
    else:
        args = ['python3', path]
    return args + (['--check'] if check else [])

def check_runtime():
    result = subprocess.run(command(True), capture_output=True, timeout=15,
                            creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    if result.returncode:
        raise ValueError('受限运行环境未就绪，请按部署指南安装 WSL Ubuntu 和 readonly-v1 环境')
    return json.loads(result.stdout)

def validate_requirements(text):
    for line in text.splitlines():
        line = line.split('#', 1)[0].strip()
        if not line:
            continue
        match = re.fullmatch(r'([A-Za-z0-9_-]+)(?:==([0-9.]+))?', line)
        if not match or match[1].lower() not in SUPPORTED or (match[2] and match[2] != SUPPORTED[match[1].lower()]):
            raise ValueError('受限环境支持标准库、requests==2.32.5、openpyxl==3.1.5、beautifulsoup4==4.13.5；本任务其他依赖暂不能自动试跑，原版本保留')

def validate_contract(value):
    if not isinstance(value, dict) or set(value) - {'file', 'format', 'min_rows', 'max_rows', 'required', 'unique_by', 'equals', 'positive', 'urls', 'effects'}:
        raise ValueError('验收条件字段无效')
    result = {'file': '', 'format': 'csv', 'min_rows': 1, 'max_rows': 100000, 'required': [], 'unique_by': [], 'equals': {}, 'positive': [], 'urls': [], **value}
    if not safe_name(result['file']):
        raise ValueError('结果文件名无效，请使用单个文件名')
    if result['format'] not in {'csv', 'json', 'xlsx'}:
        raise ValueError('结果格式应为 csv/json/xlsx')
    if result.get('effects','') not in {'','artifacts_only'}:
        raise ValueError('目前只支持声明仅生成产物的任务自动启用修复版本')
    if any(type(result[key]) is not int for key in ('min_rows', 'max_rows')) or not 0 <= result['min_rows'] <= result['max_rows'] <= 100000:
        raise ValueError('结果行数范围无效')
    for key in ('required', 'unique_by', 'positive', 'urls'):
        if not isinstance(result[key], list) or len(result[key]) > 20 or any(not isinstance(item, str) or not item or len(item) > 2000 for item in result[key]):
            raise ValueError('列名或网址列表无效')
    if not result['required'] or not isinstance(result['equals'], dict) or len(result['equals']) > 20:
        raise ValueError('至少填写一个必需字段，固定值条件必须是对象')
    if any(not isinstance(k, str) or not isinstance(v, (str, int, float, bool)) for k, v in result['equals'].items()):
        raise ValueError('固定值条件无效')
    if not set(result['unique_by'] + result['positive'] + list(result['equals'])).issubset(result['required']):
        raise ValueError('唯一值、正数及固定值字段必须同时列入必需字段')
    for url in result['urls']:
        parsed = urlsplit(url)
        if parsed.scheme not in {'http','https'} or not parsed.hostname or parsed.username or parsed.password or parsed.fragment or parsed.port not in (None,80,443):
            raise ValueError('网页快照需填写完整的公开 HTTP/HTTPS 地址')
    return result

def capture_pages(contract):
    hosts = {urlsplit(url).hostname.lower().rstrip('.') for url in contract.get('urls', [])}
    pages = {}
    total = 0
    for url in contract.get('urls', []):
        page = ai_tools.fetch_page(url, hosts, raw=True)
        if page.get('status') != 200 or 'body' not in page:
            raise ValueError(f"网页快照不可用（HTTP {page.get('status')}），需检查访问权限；未触发代码修复")
        total += len(page['body'])
        if total > 8 * 1024 * 1024:
            raise ValueError('本次网页快照超过总量上限')
        pages[url] = page
    return pages

async def execute_source(source, *, pages, template='', timeout=120, stop=None):
    payload = json.dumps({'source': source, 'pages': pages, 'template': template, 'timeout': timeout}, ensure_ascii=True).encode()
    process = await asyncio.create_subprocess_exec(*command(), stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    async def drain(stream):
        data = bytearray()
        while chunk := await stream.read(65536):
            data.extend(chunk)
            if len(data) > MAX_OUTPUT:
                raise ValueError('受限执行输出过大')
        return bytes(data)
    async def communicate():
        # Drain concurrently so bounded pipes cannot deadlock on large input/output.
        out, err = asyncio.create_task(drain(process.stdout)), asyncio.create_task(drain(process.stderr))
        try:
            process.stdin.write(payload)
            await process.stdin.drain()
            process.stdin.close()
            result = await asyncio.gather(out, err)
            await process.wait()
            return result
        finally:
            for task in (out, err):
                if not task.done(): task.cancel()
            await asyncio.gather(out, err, return_exceptions=True)
    running = asyncio.create_task(communicate())
    halted = asyncio.create_task(stop.wait()) if stop else None
    try:
        done, _ = await asyncio.wait([running] + ([halted] if halted else []), timeout=timeout + 20, return_when=asyncio.FIRST_COMPLETED)
        if halted and halted in done:
            raise InterruptedError('已请求停止受限执行')
        if running not in done:
            raise TimeoutError('受限执行达到时间预算')
        stdout, stderr = running.result()
        if process.returncode:
            raise ValueError('受限执行未完成：' + stderr.decode('utf-8', errors='replace')[-1000:])
        return decode_result(json.loads(stdout))
    finally:
        if process.returncode is None:
            # WSL driver has its own wall deadline; stopping the Windows bridge alone
            # does not prove Linux exit. Wait for that deadline before releasing queue.
            try: await asyncio.wait_for(asyncio.shield(running), timeout=timeout + 10)
            except Exception:
                if process.returncode is None: process.kill()
                await process.wait()
        if halted: halted.cancel()
        if not running.done(): running.cancel()
        await asyncio.gather(running, *([halted] if halted else []), return_exceptions=True)

def decode_result(reply):
    if not isinstance(reply, dict) or type(reply.get('exit_code')) is not int or not isinstance(reply.get('artifacts'), list) or len(reply['artifacts']) > 20:
        raise ValueError('受限执行回执无效')
    files = {}
    for item in reply['artifacts']:
        name = item['name']
        if not safe_name(name) or name.casefold() in {key.casefold() for key in files}:
            raise ValueError('产物文件名无效或重复')
        files[name] = base64.b64decode(item['data'], validate=True)
    if sum(map(len, files.values())) > 8 * 1024 * 1024:
        raise ValueError('产物超过限额')
    return {'exit_code': reply['exit_code'], 'log': str(reply.get('log', ''))[-64000:], 'files': files,
            'result': str(reply.get('result', ''))[:65536]}

def validate_result(result, contract):
    if result['exit_code'] != 0:
        raise ValueError(result['log'][-5000:] or '脚本执行失败')
    if result.get('result'):
        receipt = json.loads(result['result'])
        if not isinstance(receipt,dict) or receipt.get('outcome')!='success':
            raise ValueError('脚本业务回执没有报告成功，保留原版本')
    if not contract:
        return {'passed': True, 'message': '运行验证通过'}
    raw = result['files'].get(contract['file'])
    if raw is None:
        raise ValueError('缺少约定的结果文件：' + contract['file'])
    if contract['format'] == 'csv':
        rows = csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))
    elif contract['format'] == 'json':
        rows = json.loads(raw)
        if not isinstance(rows,list): raise ValueError('JSON 结果必须是记录数组')
    else:
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist())>1000 or sum(item.file_size for item in archive.infolist())>32*1024*1024:
                raise ValueError('Excel 解压后的内容超过验收上限')
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        if (workbook.active.max_column or 0)>200:
            workbook.close()
            raise ValueError('Excel 结果超过 200 列验收上限')
        workbook.active.reset_dimensions()
        values = workbook.active.iter_rows(values_only=True,max_col=200)
        header = [str(v) if v is not None else '' for v in next(values, [])]
        def excel_rows():
            try:
                for row in values:
                    yield dict(zip(header,row))
            finally: workbook.close()
        rows = excel_rows()
    seen = set()
    count = 0
    try:
        for row in rows:
            count += 1
            if count>contract['max_rows']: raise ValueError('结果行数超过原验收上限')
            validate_row(row,contract,seen)
    finally:
        if hasattr(rows,'close'): rows.close()
    if count<contract['min_rows']:
        raise ValueError(f"结果行数未通过原验收条件（要求 {contract['min_rows']}–{contract['max_rows']}）")
    return {'passed': True, 'rows': count, 'file': contract['file'], 'message': f'原验收条件通过，实际 {count} 条'}


def validate_repair_structure(original, candidate):
    """Reject deletion/exception suppression as a substitute for a minimal repair.

    This is a regression guard, not a claim to infer arbitrary business rules.
    Runtime verification must independently reproduce the failure with frozen inputs.
    """
    new = ast.parse(candidate)
    if not new.body or SequenceMatcher(None, original, candidate, autojunk=False).ratio() < .6:
        raise ValueError('候选改动过大，未通过最小修复检查')
    try:
        old = ast.parse(original)
    except SyntaxError:
        # Keep a conservative textual guard for code that cannot yet be parsed.
        for marker in ('def ', 'return ', 'assert ', 'import '):
            if candidate.count(marker) < original.count(marker):
                raise ValueError('候选删除了原有逻辑，未通过最小修复检查')
        if any(isinstance(node, (ast.Try, ast.TryStar)) for node in ast.walk(new)):
            raise ValueError('不能通过捕获异常掩盖原始错误')
        return

    def required_nodes(tree):
        nodes = Counter()
        for node in ast.walk(tree):
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                nodes[(type(node).__name__, node.name)] += 1
            elif isinstance(node, (ast.Import, ast.ImportFrom)):
                nodes[('import', ast.dump(node))] += 1
            elif isinstance(node, (ast.Return, ast.Assert)):
                nodes[(type(node).__name__, bool(getattr(node, 'value', True)))] += 1
            elif isinstance(node, ast.Call):
                nodes[('call', ast.dump(node.func))] += 1
        return nodes

    if required_nodes(old) - required_nodes(new):
        raise ValueError('候选删除了原有函数、调用或校验，未通过最小修复检查')
    before = sum(isinstance(node, (ast.Try, ast.TryStar)) for node in ast.walk(old))
    after = sum(isinstance(node, (ast.Try, ast.TryStar)) for node in ast.walk(new))
    if after > before:
        raise ValueError('不能通过捕获异常掩盖原始错误')
    for node in ast.walk(new):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in {'exit', 'quit'}:
            raise ValueError('不能通过提前退出跳过任务')

def validate_row(row, contract, seen):
    if not isinstance(row, dict) or any(key not in row or row[key] is None or str(row[key]).strip() == '' for key in contract['required']):
        raise ValueError('结果存在缺失字段或空值')
    for key, value in contract['equals'].items():
        if str(row[key]) != str(value): raise ValueError('结果未满足固定值条件：' + key)
    for key in contract['positive']:
        try:
            number = float(row[key])
            if not 0 < number < float('inf'): raise ValueError()
        except (ValueError, TypeError): raise ValueError('结果未满足正数条件：' + key) from None
    if contract['unique_by']:
        identity = tuple(str(row[key]) for key in contract['unique_by'])
        if identity in seen: raise ValueError('结果存在重复记录')
        seen.add(identity)
